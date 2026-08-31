from openpyxl import load_workbook


class ExcelUtils:

    @staticmethod
    def read(file_path, sheet_name=None):

        workbook = load_workbook(
            file_path,
            data_only=True
        )

        sheet = (
            workbook[sheet_name]
            if sheet_name
            else workbook.active
        )

        rows = list(
            sheet.iter_rows(
                values_only=True
            )
        )

        if not rows:
            return []

        headers = [
            str(value).strip()
            if value is not None
            else ""
            for value in rows[0]
        ]

        result = []

        for row in rows[1:]:

            # Skip empty row
            if all(
                value is None
                for value in row
            ):
                continue

            item = {}

            for index, header in enumerate(headers):

                if not header:
                    continue

                value = (
                    row[index]
                    if index < len(row)
                    else None
                )

                item[header] = value

            result.append(item)

        return result